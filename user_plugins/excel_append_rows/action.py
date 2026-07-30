"""把工作流中的字典记录写入 Excel 的通用用户插件。"""
from __future__ import annotations

from pathlib import Path
import threading
from typing import Any


_WRITE_LOCK = threading.Lock()


class ExcelAppendError(RuntimeError):
    pass


def _records(value: Any) -> list[dict[str, Any]]:
    if value in (None, ""):
        return []
    if not isinstance(value, list) or not all(isinstance(item, dict) for item in value):
        raise ExcelAppendError("“要追加的记录”必须引用一个字典列表，例如 {{ steps.upload.result.uploaded_files }}")
    return value


def _headers(records: list[dict[str, Any]], existing: list[Any]) -> list[str]:
    if existing:
        return [str(value) if value is not None else "" for value in existing]
    ordered: list[str] = []
    for record in records:
        for key in record:
            if key not in ordered:
                ordered.append(key)
    return ordered


def run_with_context(action_info: dict[str, Any], params: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    raw_path = str(params.get("excel_path", "")).strip()
    if not raw_path:
        raise ExcelAppendError("请填写 Excel 文件")
    path = Path(raw_path).expanduser().resolve()
    if path.suffix.lower() != ".xlsx":
        raise ExcelAppendError("Excel 文件必须以 .xlsx 结尾")
    records = _records(params.get("records"))
    if not records:
        return {"appended": 0, "skipped": 0, "workbook_path": str(path)}
    sheet_name = str(params.get("sheet_name", "归档记录")).strip() or "归档记录"
    unique_columns = [item.strip() for item in str(params.get("deduplicate_by", "")).split(",") if item.strip()]
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise ExcelAppendError("此插件需要 openpyxl 才能写入 .xlsx") from exc

    with _WRITE_LOCK:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = load_workbook(path) if path.exists() else Workbook()
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
        sheet.title = sheet_name
        first_row = [cell.value for cell in sheet[1]] if sheet.max_row else []
        existing_header = first_row if any(value is not None for value in first_row) else []
        headers = _headers(records, existing_header)
        if not headers:
            return {"appended": 0, "skipped": 0, "workbook_path": str(path)}
        if not existing_header:
            sheet.append(headers)
            sheet.freeze_panes = "A2"
        elif headers != [str(value) if value is not None else "" for value in existing_header]:
            missing = [name for name in headers if name not in existing_header]
            if missing:
                raise ExcelAppendError(f"现有工作表缺少列：{', '.join(missing)}")
            headers = [str(value) if value is not None else "" for value in existing_header]
        indexes = {name: index for index, name in enumerate(headers)}
        missing_keys = [name for name in unique_columns if name not in indexes]
        if missing_keys:
            raise ExcelAppendError(f"去重列不在记录中：{', '.join(missing_keys)}")
        seen = set()
        if unique_columns:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                seen.add(tuple(row[indexes[name]] for name in unique_columns))
        appended = skipped = 0
        for record in records:
            key = tuple(record.get(name) for name in unique_columns)
            if unique_columns and key in seen:
                skipped += 1
                continue
            sheet.append([record.get(name, "") for name in headers])
            seen.add(key)
            appended += 1
        workbook.save(path)
    return {"appended": appended, "skipped": skipped, "workbook_path": str(path)}


def run(action_info: dict[str, Any], params: dict[str, Any]) -> dict[str, Any]:
    return run_with_context(action_info, params, {"event": {}, "steps": {}})
